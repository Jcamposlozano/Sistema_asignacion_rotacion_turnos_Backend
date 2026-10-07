from __future__ import annotations

from io import BytesIO

import openpyxl
from sqlalchemy import create_engine, delete, select
from sqlalchemy.orm import sessionmaker

from asignacion_estudiantes_service.application.analysis_service import obtener_analisis_asignacion
from asignacion_estudiantes_service.application.asignacion_service import (
    crear_ejecucion_desde_matriz,
    ejecutar_asignacion,
    limpiar_resultados_huerfanos,
    listar_ejecuciones_parametria,
    listar_resultados,
    obtener_ultima_ejecucion_parametria,
)
from asignacion_estudiantes_service.application.diagnostic_service import (
    obtener_diagnostico_asignacion,
)
from asignacion_estudiantes_service.application.dashboard_service import obtener_resumen_gerencial
from asignacion_estudiantes_service.application.file_import_service import (
    leer_estudiantes_archivo,
    leer_instituciones_archivo,
    leer_parametria_plantilla,
)
from asignacion_estudiantes_service.application.parametria_service import (
    actualizar_parametria,
    crear_parametria,
    eliminar_parametria,
    obtener_parametria_completa,
)
from asignacion_estudiantes_service.application.optimization_service import (
    optimizar_cupos_parametria,
    optimizar_y_validar_parametria,
)
from asignacion_estudiantes_service.domain.models import (
    Asignacion,
    AsignacionEjecucion,
    Base,
    Estudiante,
    Parametria,
)
from asignacion_estudiantes_service.domain.schemas import (
    MatrizAsignacionIn,
    OptimizarParametriaIn,
    ParametriaCargaIn,
)


def test_carga_parametria_y_ejecuta_asignacion(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path}/test.db", future=True)
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)

    payload = ParametriaCargaIn(
        nombre="Escenario test",
        instituciones=[
            {"institucion": "CLINICA DE TENJO", "especialidad": "MEDICINA FAMILIAR", "cantidad": 5},
            {
                "institucion": "CLINICA UNIVERSIDAD DE LA SABANA",
                "especialidad": "MEDICINA FAMILIAR",
                "cantidad": 6,
            },
            {
                "institucion": "CLINICA UNIVERSIDAD DE LA SABANA",
                "especialidad": "MEDICINA INTERNA",
                "cantidad": 12,
            },
            {
                "institucion": "CLINICA UNIVERSIDAD DE LA SABANA",
                "especialidad": "NEUROLOGIA",
                "cantidad": 3,
            },
            {"institucion": "HOSPITAL SIMULADO", "especialidad": "SALUD MENTAL", "cantidad": 18},
            {
                "institucion": "HOSPITAL SIMULADO",
                "especialidad": "MEDICINA INTERNA HOSPITAL SIMULADO",
                "cantidad": 20,
            },
        ],
        restricciones=[
            {
                "estudiante": "352043",
                "institucion": "CLINICA DE TENJO",
                "especialidad": "MEDICINA FAMILIAR",
            }
        ],
        asignacion={
            "numero_periodos": 8,
            "estudiantes_aleatorio": False,
            "instituciones_aleatorio": False,
            "asignacion": "CONCENTRADA",
            "institucion_residual": "CLINICA UNIVERSIDAD DE LA SABANA",
        },
        distribucion_periodos=[
            {"especialidad": "MEDICINA FAMILIAR", "asignacion": 1},
            {"especialidad": "MEDICINA INTERNA", "asignacion": 4},
            {"especialidad": "NEUROLOGIA", "asignacion": 1},
            {"especialidad": "SALUD MENTAL", "asignacion": 1},
            {"especialidad": "MEDICINA INTERNA HOSPITAL SIMULADO", "asignacion": 1},
        ],
        combinaciones=[
            {
                "especialidad_1": "MEDICINA INTERNA",
                "especialidad_2": "MEDICINA INTERNA",
                "bloque": "X4",
            },
            {
                "especialidad_1": "MEDICINA INTERNA HOSPITAL SIMULADO",
                "especialidad_2": "SALUD MENTAL",
                "bloque": "X2",
            },
            {
                "especialidad_1": "MEDICINA FAMILIAR",
                "especialidad_2": "NEUROLOGIA",
                "bloque": "X2",
            },
        ],
        estudiantes=[
            {"id": "352043", "nombre": "Estudiante Uno", "semestre": "6"},
            {"id": "350248", "nombre": "Estudiante Dos", "semestre": "6"},
        ],
    )

    with SessionLocal() as session:
        parametria = crear_parametria(session, payload)
        ejecucion = ejecutar_asignacion(session, parametria.id)
        estudiante = session.get(Estudiante, "352043")
        analisis = obtener_analisis_asignacion(session, ejecucion.id)

    assert ejecucion.estado == "COMPLETADA"
    assert ejecucion.mensaje is not None
    assert "intento" in ejecucion.mensaje
    assert ejecucion.total_estudiantes == 2
    assert ejecucion.total_asignaciones == 16
    assert ejecucion.total_pendientes == 0
    assert estudiante is not None
    assert estudiante.semestre == "6"
    assert analisis is not None
    assert analisis["resumen"]["asignaciones_realizadas"] == 16
    assert len(analisis["resumen_especialidad"]) == 5


def test_importa_estudiantes_desde_excel_xlsx():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["ID", "NOMBRE", "SEMESTRE"])
    sheet.append(["352043", "Estudiante Uno", "6"])
    stream = BytesIO()
    workbook.save(stream)

    estudiantes = leer_estudiantes_archivo("estudiantes.xlsx", stream.getvalue())

    assert len(estudiantes) == 1
    assert estudiantes[0].id == "352043"
    assert estudiantes[0].nombre == "Estudiante Uno"
    assert estudiantes[0].semestre == "6"


def test_importa_instituciones_desde_excel_xlsx():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["INSTITUCION", "ESPECIALIDAD", "CANTIDAD"])
    sheet.append(["CLINICA A", "MEDICINA INTERNA", 6])
    stream = BytesIO()
    workbook.save(stream)

    instituciones = leer_instituciones_archivo("instituciones.xlsx", stream.getvalue())

    assert len(instituciones) == 1
    assert instituciones[0].institucion == "CLINICA A"
    assert instituciones[0].especialidad == "MEDICINA INTERNA"
    assert instituciones[0].cantidad == 6


def test_importa_plantilla_parametria_desde_excel_xlsx():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Asignacion"
    sheet.append(["Campo", "Valor"])
    sheet.append(["Nombre", "Escenario plantilla"])
    sheet.append(["Numero Periodos", 6])
    sheet.append(["Estudiantes Aleatorio", "False"])
    sheet.append(["Instituciones Aleatorio", "False"])
    sheet.append(["Asignacion", "BALANCEADA"])
    sheet.append(["Institucion residual", "CLINICA A"])

    instituciones = workbook.create_sheet("Instituciones")
    instituciones.append(["INSTITUCION", "ESPECIALIDAD", "CANTIDAD"])
    instituciones.append(["CLINICA A", "MEDICINA INTERNA", 6])

    estudiantes = workbook.create_sheet("Estudiantes")
    estudiantes.append(["ID", "NOMBRE", "SEMESTRE"])
    estudiantes.append(["1", "Estudiante Uno", "6"])

    distribucion = workbook.create_sheet("Distribucion")
    distribucion.append(["ESPECIALIDAD", "ASIGNACION"])
    distribucion.append(["MEDICINA INTERNA", 6])

    restricciones = workbook.create_sheet("Restricciones")
    restricciones.append(["ESTUDIANTE", "INSTITUCION", "ESPECIALIDAD"])

    combinaciones = workbook.create_sheet("Combinaciones")
    combinaciones.append(["ESPECIALIDAD 1", "ESPECIALIDAD 2", "BLOQUE"])

    stream = BytesIO()
    workbook.save(stream)

    payload = leer_parametria_plantilla("plantilla.xlsx", stream.getvalue())

    assert payload.nombre == "Escenario plantilla"
    assert payload.asignacion.numero_periodos == 6
    assert payload.asignacion.estudiantes_aleatorio is False
    assert payload.instituciones[0].institucion == "CLINICA A"
    assert payload.estudiantes[0].nombre == "Estudiante Uno"
    assert payload.distribucion_periodos[0].asignacion == 6


def test_permite_parametria_de_cinco_periodos(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path}/test.db", future=True)
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)

    payload = ParametriaCargaIn(
        nombre="Cinco periodos",
        instituciones=[
            {"institucion": "CLINICA A", "especialidad": "MEDICINA INTERNA", "cantidad": 2},
        ],
        restricciones=[],
        asignacion={
            "numero_periodos": 5,
            "estudiantes_aleatorio": False,
            "instituciones_aleatorio": False,
            "asignacion": "BALANCEADA",
            "institucion_residual": "CLINICA A",
        },
        distribucion_periodos=[{"especialidad": "MEDICINA INTERNA", "asignacion": 5}],
        combinaciones=[],
        estudiantes=[{"id": "1", "nombre": "Estudiante Uno", "semestre": "5"}],
    )

    with SessionLocal() as session:
        parametria = crear_parametria(session, payload)
        ejecucion = ejecutar_asignacion(session, parametria.id, max_intentos=1)
        resultados = listar_resultados(session, ejecucion.id)

    assert ejecucion.total_asignaciones == 5
    assert ejecucion.total_pendientes == 0
    assert {row.periodo for row in resultados} == {1, 2, 3, 4, 5}


def test_actualiza_y_elimina_parametria(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path}/test.db", future=True)
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)

    payload = ParametriaCargaIn(
        nombre="Escenario original",
        instituciones=[
            {"institucion": "CLINICA DE TENJO", "especialidad": "MEDICINA FAMILIAR", "cantidad": 2}
        ],
        restricciones=[],
        asignacion={
            "numero_periodos": 6,
            "estudiantes_aleatorio": False,
            "instituciones_aleatorio": False,
            "asignacion": "BALANCEADA",
            "institucion_residual": "CLINICA UNIVERSIDAD DE LA SABANA",
        },
        distribucion_periodos=[{"especialidad": "MEDICINA FAMILIAR", "asignacion": 6}],
        combinaciones=[],
        estudiantes=[{"id": "352043", "nombre": "Estudiante Uno", "semestre": "6"}],
    )
    payload_actualizado = ParametriaCargaIn(
        **{
            **payload.model_dump(),
            "nombre": "Escenario actualizado",
            "estudiantes": [{"id": "352043", "nombre": "Estudiante Uno", "semestre": "7"}],
        }
    )

    with SessionLocal() as session:
        parametria = crear_parametria(session, payload)
        ejecucion = ejecutar_asignacion(session, parametria.id)

        actualizada = actualizar_parametria(session, parametria.id, payload_actualizado)
        estudiante = session.get(Estudiante, "352043")

        assert ejecucion.id > 0
        assert actualizada is not None
        assert actualizada.id == parametria.id
        assert actualizada.nombre == "Escenario actualizado"
        assert estudiante is not None
        assert estudiante.semestre == "7"

        eliminada = eliminar_parametria(session, parametria.id)
        assert eliminada is True
        assert obtener_parametria_completa(session, parametria.id) is None
        assert obtener_ultima_ejecucion_parametria(session, parametria.id) is None
        assert list(session.scalars(select(Asignacion).where(Asignacion.parametria_id == parametria.id))) == []


def test_restriccion_especifica_es_obligatoria_sin_fallback(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path}/test.db", future=True)
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)

    payload = ParametriaCargaIn(
        nombre="Restriccion especifica obligatoria",
        instituciones=[
            {"institucion": "CLINICA A", "especialidad": "MEDICINA INTERNA", "cantidad": 0},
            {"institucion": "CLINICA B", "especialidad": "MEDICINA INTERNA", "cantidad": 10},
        ],
        restricciones=[
            {
                "estudiante": "1",
                "institucion": "CLINICA A",
                "especialidad": "MEDICINA INTERNA",
            }
        ],
        asignacion={
            "numero_periodos": 6,
            "estudiantes_aleatorio": False,
            "instituciones_aleatorio": False,
            "asignacion": "BALANCEADA",
            "institucion_residual": "CLINICA B",
        },
        distribucion_periodos=[{"especialidad": "MEDICINA INTERNA", "asignacion": 6}],
        combinaciones=[],
        estudiantes=[{"id": "1", "nombre": "Estudiante Uno", "semestre": "6"}],
    )

    with SessionLocal() as session:
        parametria = crear_parametria(session, payload)
        ejecucion = ejecutar_asignacion(session, parametria.id)
        resultados = listar_resultados(session, ejecucion.id)

    assert ejecucion.total_pendientes == 6
    assert {row.institucion for row in resultados} == {"PENDIENTE"}


def test_restriccion_general_prioriza_solo_especialidades_disponibles(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path}/test.db", future=True)
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)

    payload = ParametriaCargaIn(
        nombre="Restriccion general preferente",
        instituciones=[
            {"institucion": "CLINICA A", "especialidad": "MEDICINA INTERNA", "cantidad": 10},
            {"institucion": "CLINICA B", "especialidad": "NEUROLOGIA", "cantidad": 10},
        ],
        restricciones=[{"estudiante": "1", "institucion": "CLINICA A"}],
        asignacion={
            "numero_periodos": 6,
            "estudiantes_aleatorio": False,
            "instituciones_aleatorio": False,
            "asignacion": "BALANCEADA",
            "institucion_residual": "CLINICA B",
        },
        distribucion_periodos=[
            {"especialidad": "MEDICINA INTERNA", "asignacion": 3},
            {"especialidad": "NEUROLOGIA", "asignacion": 3},
        ],
        combinaciones=[],
        estudiantes=[{"id": "1", "nombre": "Estudiante Uno", "semestre": "6"}],
    )

    with SessionLocal() as session:
        parametria = crear_parametria(session, payload)
        ejecucion = ejecutar_asignacion(session, parametria.id)
        resultados = listar_resultados(session, ejecucion.id)

    assert ejecucion.total_pendientes == 0
    assert {row.institucion for row in resultados if row.especialidad == "MEDICINA INTERNA"} == {
        "CLINICA A"
    }
    assert {row.institucion for row in resultados if row.especialidad == "NEUROLOGIA"} == {
        "CLINICA B"
    }


def test_estudiantes_con_restriccion_se_asignan_primero(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path}/test.db", future=True)
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)

    payload = ParametriaCargaIn(
        nombre="Prioridad de estudiantes restringidos",
        instituciones=[
            {"institucion": "CLINICA A", "especialidad": "MEDICINA INTERNA", "cantidad": 1},
        ],
        restricciones=[
            {
                "estudiante": "RESTRINGIDO",
                "institucion": "CLINICA A",
                "especialidad": "MEDICINA INTERNA",
            }
        ],
        asignacion={
            "numero_periodos": 6,
            "estudiantes_aleatorio": False,
            "instituciones_aleatorio": False,
            "asignacion": "BALANCEADA",
            "institucion_residual": "CLINICA A",
        },
        distribucion_periodos=[{"especialidad": "MEDICINA INTERNA", "asignacion": 6}],
        combinaciones=[],
        estudiantes=[
            {"id": "LIBRE", "nombre": "Sin restriccion", "semestre": "6"},
            {"id": "RESTRINGIDO", "nombre": "Con restriccion", "semestre": "6"},
        ],
    )

    with SessionLocal() as session:
        parametria = crear_parametria(session, payload)
        ejecucion = ejecutar_asignacion(session, parametria.id)
        resultados = listar_resultados(session, ejecucion.id)

    restringido_periodo_1 = next(
        row for row in resultados if row.estudiante_id == "RESTRINGIDO" and row.periodo == 1
    )
    libre_periodo_1 = next(
        row for row in resultados if row.estudiante_id == "LIBRE" and row.periodo == 1
    )

    assert restringido_periodo_1.institucion == "CLINICA A"
    assert libre_periodo_1.institucion == "PENDIENTE"


def test_ejecucion_respeta_maximo_de_intentos_configurado(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path}/test.db", future=True)
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)

    payload = ParametriaCargaIn(
        nombre="Maximo intentos",
        instituciones=[
            {"institucion": "CLINICA A", "especialidad": "MEDICINA INTERNA", "cantidad": 2},
        ],
        restricciones=[],
        asignacion={
            "numero_periodos": 6,
            "estudiantes_aleatorio": True,
            "instituciones_aleatorio": True,
            "asignacion": "BALANCEADA",
            "institucion_residual": "CLINICA A",
        },
        distribucion_periodos=[{"especialidad": "MEDICINA INTERNA", "asignacion": 6}],
        combinaciones=[],
        estudiantes=[{"id": "1", "nombre": "Estudiante Uno", "semestre": "6"}],
    )

    with SessionLocal() as session:
        parametria = crear_parametria(session, payload)
        ejecucion = ejecutar_asignacion(session, parametria.id, max_intentos=1)

    assert ejecucion.mensaje is not None
    assert "entre 1 intento" in ejecucion.mensaje


def test_obtiene_ultima_ejecucion_guardada_de_parametria(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path}/test.db", future=True)
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)

    payload = ParametriaCargaIn(
        nombre="Ultima ejecucion",
        instituciones=[
            {"institucion": "CLINICA A", "especialidad": "MEDICINA INTERNA", "cantidad": 2},
        ],
        restricciones=[],
        asignacion={
            "numero_periodos": 6,
            "estudiantes_aleatorio": False,
            "instituciones_aleatorio": False,
            "asignacion": "BALANCEADA",
            "institucion_residual": "CLINICA A",
        },
        distribucion_periodos=[{"especialidad": "MEDICINA INTERNA", "asignacion": 6}],
        combinaciones=[],
        estudiantes=[{"id": "1", "nombre": "Estudiante Uno", "semestre": "6"}],
    )

    with SessionLocal() as session:
        parametria = crear_parametria(session, payload)
        primera = ejecutar_asignacion(session, parametria.id, max_intentos=1)
        segunda = ejecutar_asignacion(session, parametria.id, max_intentos=1)
        ultima = obtener_ultima_ejecucion_parametria(session, parametria.id)
        primera_id = primera.id
        segunda_id = segunda.id
        ultima_id = ultima.id if ultima else None

    assert ultima is not None
    assert primera_id < segunda_id
    assert ultima_id == segunda_id


def test_carga_matriz_ajustada_como_ejecucion_manual(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path}/test.db", future=True)
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)

    payload = ParametriaCargaIn(
        nombre="Matriz ajustada",
        instituciones=[
            {"institucion": "CLINICA A", "especialidad": "MEDICINA INTERNA", "cantidad": 6},
        ],
        restricciones=[],
        asignacion={
            "numero_periodos": 6,
            "estudiantes_aleatorio": False,
            "instituciones_aleatorio": False,
            "asignacion": "BALANCEADA",
            "institucion_residual": "CLINICA A",
        },
        distribucion_periodos=[{"especialidad": "MEDICINA INTERNA", "asignacion": 6}],
        combinaciones=[],
        estudiantes=[{"id": "1", "nombre": "Estudiante Uno", "semestre": "6"}],
    )

    matriz = MatrizAsignacionIn(
        filas=[
            {
                "estudiante_id": "1",
                "periodo": periodo,
                "especialidad": "MEDICINA INTERNA",
                "institucion": "CLINICA A",
            }
            for periodo in range(1, 7)
        ]
    )

    with SessionLocal() as session:
        parametria = crear_parametria(session, payload)
        ejecucion = crear_ejecucion_desde_matriz(session, parametria.id, matriz)
        resultados = listar_resultados(session, ejecucion.id)
        ejecuciones = listar_ejecuciones_parametria(session, parametria.id)
        ultima = obtener_ultima_ejecucion_parametria(session, parametria.id)

    assert ejecucion.estado == "AJUSTADA_MANUALMENTE"
    assert ejecucion.total_asignaciones == 6
    assert ejecucion.total_pendientes == 0
    assert resultados[0].estudiante_nombre == "Estudiante Uno"
    assert ejecuciones[0].id == ejecucion.id
    assert ultima is not None
    assert ultima.id == ejecucion.id


def test_limpia_resultados_huerfanos_para_dashboard(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path}/test.db", future=True)
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)

    payload = ParametriaCargaIn(
        nombre="Huerfanos dashboard",
        instituciones=[
            {"institucion": "CLINICA A", "especialidad": "MEDICINA INTERNA", "cantidad": 2},
        ],
        restricciones=[],
        asignacion={
            "numero_periodos": 6,
            "estudiantes_aleatorio": False,
            "instituciones_aleatorio": False,
            "asignacion": "BALANCEADA",
            "institucion_residual": "CLINICA A",
        },
        distribucion_periodos=[{"especialidad": "MEDICINA INTERNA", "asignacion": 6}],
        combinaciones=[],
        estudiantes=[{"id": "1", "nombre": "Estudiante Uno", "semestre": "6"}],
    )

    with SessionLocal() as session:
        parametria = crear_parametria(session, payload)
        ejecucion = ejecutar_asignacion(session, parametria.id, max_intentos=1)
        parametria_id = parametria.id
        ejecucion_id = ejecucion.id
        session.execute(delete(Parametria).where(Parametria.id == parametria_id))
        session.commit()

        limpieza = limpiar_resultados_huerfanos(session)

        assert limpieza["ejecuciones_eliminadas"] == 1
        assert limpieza["asignaciones_eliminadas"] >= 1
        assert limpieza["detalles_parametria_eliminados"] >= 1
        assert session.get(AsignacionEjecucion, ejecucion_id) is None
        assert list(session.scalars(select(Asignacion).where(Asignacion.parametria_id == parametria_id))) == []


def test_diagnostico_detecta_pendientes_y_recomienda_acciones(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path}/test.db", future=True)
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)

    payload = ParametriaCargaIn(
        nombre="Diagnostico con pendientes",
        instituciones=[
            {"institucion": "CLINICA A", "especialidad": "MEDICINA INTERNA", "cantidad": 0},
            {"institucion": "CLINICA B", "especialidad": "MEDICINA INTERNA", "cantidad": 10},
        ],
        restricciones=[
            {
                "estudiante": "1",
                "institucion": "CLINICA A",
                "especialidad": "MEDICINA INTERNA",
            }
        ],
        asignacion={
            "numero_periodos": 6,
            "estudiantes_aleatorio": False,
            "instituciones_aleatorio": False,
            "asignacion": "BALANCEADA",
            "institucion_residual": "CLINICA B",
        },
        distribucion_periodos=[{"especialidad": "MEDICINA INTERNA", "asignacion": 6}],
        combinaciones=[],
        estudiantes=[{"id": "1", "nombre": "Estudiante Uno", "semestre": "6"}],
    )

    with SessionLocal() as session:
        parametria = crear_parametria(session, payload)
        ejecucion = ejecutar_asignacion(session, parametria.id)
        diagnostico = obtener_diagnostico_asignacion(session, ejecucion.id)

    assert diagnostico is not None
    assert diagnostico["nivel"] == "CRITICO"
    assert diagnostico["metricas"]["asignaciones_pendientes"] == 6
    assert any(item["titulo"] == "Asignaciones pendientes" for item in diagnostico["riesgos"])
    assert any(item["id"] == "filtrar-pendientes" for item in diagnostico["acciones_rapidas"])


def test_optimizacion_sugiere_cupos_minimos_balanceados_y_valida(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path}/test.db", future=True)
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)

    payload = ParametriaCargaIn(
        nombre="Optimizacion cupos",
        instituciones=[
            {"institucion": "CLINICA A", "especialidad": "MEDICINA INTERNA", "cantidad": 5},
            {"institucion": "CLINICA B", "especialidad": "MEDICINA INTERNA", "cantidad": 5},
        ],
        restricciones=[],
        asignacion={
            "numero_periodos": 6,
            "estudiantes_aleatorio": False,
            "instituciones_aleatorio": False,
            "asignacion": "BALANCEADA",
            "institucion_residual": "CLINICA A",
        },
        distribucion_periodos=[{"especialidad": "MEDICINA INTERNA", "asignacion": 6}],
        combinaciones=[],
        estudiantes=[
            {"id": "1", "nombre": "Estudiante Uno", "semestre": "6"},
            {"id": "2", "nombre": "Estudiante Dos", "semestre": "6"},
        ],
    )

    with SessionLocal() as session:
        parametria = crear_parametria(session, payload)
        optimizacion = optimizar_cupos_parametria(session, parametria.id)
        validacion = optimizar_y_validar_parametria(
            session, parametria.id, OptimizarParametriaIn(max_intentos=1)
        )
        temp = obtener_parametria_completa(session, validacion["parametria_simulada_id"])

    assert optimizacion["cupos_totales_actuales"] == 10
    assert optimizacion["cupos_totales_sugeridos"] == 2
    assert {item["cupo_sugerido"] for item in optimizacion["parametria_sugerida"]} == {1}
    assert validacion["pendientes_estimados"] == 0
    assert validacion["resultado"] == "VALIDADO"
    assert temp is None


def test_dashboard_gerencial_consolida_periodos_en_rotaciones(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path}/test.db", future=True)
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)

    payload = ParametriaCargaIn(
        nombre="Dashboard gerencial",
        instituciones=[
            {
                "institucion": "CLINICA UNIVERSIDAD DE LA SABANA",
                "especialidad": "MEDICINA INTERNA",
                "cantidad": 4,
            },
            {"institucion": "CLINICA A", "especialidad": "MEDICINA FAMILIAR", "cantidad": 4},
            {"institucion": "CLINICA B", "especialidad": "NEUROLOGIA", "cantidad": 4},
        ],
        restricciones=[],
        asignacion={
            "numero_periodos": 6,
            "estudiantes_aleatorio": False,
            "instituciones_aleatorio": False,
            "asignacion": "CONCENTRADA",
            "institucion_residual": "CLINICA UNIVERSIDAD DE LA SABANA",
        },
        distribucion_periodos=[
            {"especialidad": "MEDICINA INTERNA", "asignacion": 4},
            {"especialidad": "MEDICINA FAMILIAR", "asignacion": 1},
            {"especialidad": "NEUROLOGIA", "asignacion": 1},
        ],
        combinaciones=[
            {
                "especialidad_1": "MEDICINA INTERNA",
                "especialidad_2": "MEDICINA INTERNA",
                "bloque": "X4",
            },
            {
                "especialidad_1": "MEDICINA FAMILIAR",
                "especialidad_2": "NEUROLOGIA",
                "bloque": "X2",
            },
        ],
        estudiantes=[
            {"id": "1", "nombre": "Estudiante Uno", "semestre": "6"},
            {"id": "2", "nombre": "Estudiante Dos", "semestre": "6"},
        ],
    )

    with SessionLocal() as session:
        parametria = crear_parametria(session, payload)
        ejecutar_asignacion(session, parametria.id, max_intentos=1)
        dashboard = obtener_resumen_gerencial(session)

    assert dashboard["kpis"]["estudiantes_unicos"] == 2
    assert dashboard["kpis"]["rotaciones"] == 4
    assert dashboard["semestres_presion_rotacion"][0]["rotaciones"] == 4
    assert any(
        item["servicio"] == "MEDICINA INTERNA" and item["rotaciones"] == 2
        for item in dashboard["servicios_mayor_demanda"]
    )
    assert dashboard["top_escenarios_por_carga"][0]["rotaciones"] >= 2
