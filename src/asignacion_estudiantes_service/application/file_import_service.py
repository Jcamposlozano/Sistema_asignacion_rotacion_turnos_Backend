from __future__ import annotations

import csv
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

import openpyxl
import xlrd

from asignacion_estudiantes_service.domain.schemas import (
    AsignacionConfigIn,
    CombinacionIn,
    DistribucionPeriodoIn,
    EstudianteIn,
    InstitucionIn,
    MatrizAsignacionIn,
    MatrizAsignacionRowIn,
    ParametriaCargaIn,
    PeriodoFechaIn,
    RestriccionIn,
    limpiar_texto,
)


class ArchivoImportError(ValueError):
    pass


def leer_estudiantes_archivo(filename: str, content: bytes) -> list[EstudianteIn]:
    rows = _leer_tabla(filename, content)
    return leer_estudiantes_rows(rows)


def leer_instituciones_archivo(filename: str, content: bytes) -> list[InstitucionIn]:
    rows = _leer_tabla(filename, content)
    return _parse_instituciones(rows)


def leer_matriz_archivo(filename: str, content: bytes) -> MatrizAsignacionIn:
    rows = _leer_tabla(filename, content)
    if len(rows) < 2:
        raise ArchivoImportError("La matriz no tiene datos para cargar")

    header = [_normalizar_header(cell) for cell in rows[0]]
    id_index = _buscar_columna(header, ["id", "documento", "codigo"], default=0)
    period_columns = _period_columns(header)
    if not period_columns:
        raise ArchivoImportError("La matriz no tiene columnas de periodos")

    filas: list[MatrizAsignacionRowIn] = []
    for row in rows[1:]:
        estudiante_id = _cell(row, id_index)
        if not estudiante_id:
            continue
        for period, period_index in period_columns:
            scenario_index = _scenario_column(header, period, period_index)
            filas.append(
                MatrizAsignacionRowIn(
                    estudiante_id=estudiante_id,
                    periodo=period,
                    especialidad=_cell(row, period_index),
                    institucion=_cell(row, scenario_index),
                )
            )

    if not filas:
        raise ArchivoImportError("La matriz no tiene filas válidas")
    return MatrizAsignacionIn(
        filas=filas,
        mensaje="Malla cargada desde archivo ajustado por usuario.",
    )


def leer_parametria_plantilla(filename: str, content: bytes) -> ParametriaCargaIn:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".xls"}:
        raise ArchivoImportError("La plantilla completa debe ser un archivo .xlsx o .xls")

    sheets = _leer_libro(filename, content)
    asignacion_rows = _sheet(sheets, ["asignacion", "configuracion", "configuracion general"])
    instituciones_rows = _sheet(sheets, ["instituciones", "clinicas", "escenarios"])
    estudiantes_rows = _sheet(sheets, ["estudiantes"])
    distribucion_rows = _sheet(sheets, ["distribucion", "distribucion periodos", "periodos"])
    fechas_rows = _sheet(sheets, ["fechas", "fechas periodos", "periodo fechas"], required=False)
    restricciones_rows = _sheet(sheets, ["restricciones"], required=False)
    combinaciones_rows = _sheet(sheets, ["combinaciones", "bloques"], required=False)

    asignacion = _parse_asignacion(asignacion_rows)
    return ParametriaCargaIn(
        nombre=_parse_nombre(asignacion_rows),
        instituciones=_parse_instituciones(instituciones_rows),
        restricciones=_parse_restricciones(restricciones_rows),
        asignacion=asignacion,
        distribucion_periodos=_parse_distribucion(distribucion_rows),
        fechas_periodos=_parse_fechas_periodos(fechas_rows),
        combinaciones=_parse_combinaciones(combinaciones_rows),
        estudiantes=leer_estudiantes_rows(estudiantes_rows),
    )


def leer_estudiantes_rows(rows: list[list[str]]) -> list[EstudianteIn]:
    if not rows:
        raise ArchivoImportError("La hoja de estudiantes está vacía")
    header, data_rows = _separar_encabezado(rows, {"id", "nombre", "semestre", "documento", "codigo"})
    id_index = _buscar_columna(header, ["id", "documento", "codigo"], default=0)
    nombre_index = _buscar_columna(header, ["nombre", "estudiante"], default=1)
    semestre_index = _buscar_columna(header, ["semestre"], default=2)
    estudiantes = [
        EstudianteIn(
            id=_cell(row, id_index),
            nombre=_cell(row, nombre_index) or None,
            semestre=_cell(row, semestre_index) or None,
        )
        for row in data_rows
        if _cell(row, id_index)
    ]
    if not estudiantes:
        raise ArchivoImportError("No se encontraron estudiantes válidos en la plantilla")
    return estudiantes


def _leer_tabla(filename: str, content: bytes) -> list[list[str]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return _leer_csv(content)
    if suffix == ".xlsx":
        return _leer_xlsx(content)
    if suffix == ".xls":
        return _leer_xls(content)
    raise ArchivoImportError("Formato no soportado. Usa .csv, .xlsx o .xls")


def _leer_libro(filename: str, content: bytes) -> dict[str, list[list[str]]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".xlsx":
        workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
        return {
            _normalizar_sheet_name(sheet.title): _rows_from_values(sheet.iter_rows(values_only=True))
            for sheet in workbook.worksheets
        }
    if suffix == ".xls":
        workbook = xlrd.open_workbook(file_contents=content)
        return {
            _normalizar_sheet_name(sheet.name): _rows_from_values(
                [sheet.cell_value(row_index, column_index) for column_index in range(sheet.ncols)]
                for row_index in range(sheet.nrows)
            )
            for sheet in workbook.sheets()
        }
    raise ArchivoImportError("Formato no soportado. Usa .xlsx o .xls")


def _sheet(
    sheets: dict[str, list[list[str]]],
    aliases: list[str],
    required: bool = True,
) -> list[list[str]]:
    normalized_aliases = [_normalizar_sheet_name(alias) for alias in aliases]
    for alias in normalized_aliases:
        if alias in sheets:
            return sheets[alias]
    for name, rows in sheets.items():
        if any(alias in name for alias in normalized_aliases):
            return rows
    if required:
        raise ArchivoImportError(f"No se encontró la hoja {aliases[0]}")
    return []


def _normalizar_sheet_name(value: str) -> str:
    return _normalizar_header(value).replace("_", " ")


def _leer_csv(content: bytes) -> list[list[str]]:
    text = content.decode("utf-8-sig")
    sample = text.splitlines()[0] if text.splitlines() else ""
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    return [
        [str(cell).strip() for cell in row]
        for row in csv.reader(StringIO(text), delimiter=delimiter)
        if any(str(cell).strip() for cell in row)
    ]


def _leer_xlsx(content: bytes) -> list[list[str]]:
    workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    return _rows_from_values(sheet.iter_rows(values_only=True))


def _leer_xls(content: bytes) -> list[list[str]]:
    workbook = xlrd.open_workbook(file_contents=content)
    sheet = workbook.sheet_by_index(0)
    return _rows_from_values(
        [sheet.cell_value(row_index, column_index) for column_index in range(sheet.ncols)]
        for row_index in range(sheet.nrows)
    )


def _parse_nombre(rows: list[list[str]]) -> str:
    data = _key_value_rows(rows)
    return data.get("nombre") or data.get("parametria") or "Parametria"


def _parse_asignacion(rows: list[list[str]]) -> AsignacionConfigIn:
    data = _key_value_rows(rows)
    return AsignacionConfigIn(
        numero_periodos=_to_int(data.get("numero periodos") or data.get("periodos") or "8"),
        estudiantes_aleatorio=_to_bool(data.get("estudiantes aleatorio"), default=True),
        instituciones_aleatorio=_to_bool(data.get("instituciones aleatorio"), default=True),
        asignacion=data.get("asignacion") or data.get("estrategia") or "CONCENTRADA",
        institucion_residual=data.get("institucion residual") or data.get("residual") or "",
    )


def _parse_instituciones(rows: list[list[str]]) -> list[InstitucionIn]:
    header, data_rows = _header_rows(rows, {"institucion", "especialidad", "cantidad", "cupo"})
    institucion_index = _buscar_columna(header, ["institucion", "clinica", "escenario"], default=0)
    especialidad_index = _buscar_columna(header, ["especialidad", "servicio"], default=1)
    cantidad_index = _buscar_columna(header, ["cantidad", "cupo", "cupos"], default=2)
    instituciones = [
        InstitucionIn(
            institucion=_cell(row, institucion_index),
            especialidad=_cell(row, especialidad_index),
            cantidad=_to_int(_cell(row, cantidad_index)),
        )
        for row in data_rows
        if _cell(row, institucion_index) and _cell(row, especialidad_index)
    ]
    if not instituciones:
        raise ArchivoImportError("No se encontraron instituciones válidas en la plantilla")
    return instituciones


def _parse_restricciones(rows: list[list[str]]) -> list[RestriccionIn]:
    if not rows:
        return []
    header, data_rows = _header_rows(rows, {"estudiante", "institucion", "especialidad"})
    estudiante_index = _buscar_columna(header, ["estudiante", "id", "documento", "codigo"], default=0)
    institucion_index = _buscar_columna(header, ["institucion", "clinica", "escenario"], default=1)
    especialidad_index = _buscar_columna(header, ["especialidad", "servicio"], default=2)
    return [
        RestriccionIn(
            estudiante=_cell(row, estudiante_index),
            institucion=_cell(row, institucion_index),
            especialidad=_cell(row, especialidad_index) or None,
        )
        for row in data_rows
        if _cell(row, estudiante_index) and _cell(row, institucion_index)
    ]


def _parse_distribucion(rows: list[list[str]]) -> list[DistribucionPeriodoIn]:
    header, data_rows = _header_rows(rows, {"especialidad", "asignacion", "periodos"})
    especialidad_index = _buscar_columna(header, ["especialidad", "servicio"], default=0)
    asignacion_index = _buscar_columna(header, ["asignacion", "periodos", "cantidad"], default=1)
    distribucion = [
        DistribucionPeriodoIn(
            especialidad=_cell(row, especialidad_index),
            asignacion=_to_int(_cell(row, asignacion_index)),
        )
        for row in data_rows
        if _cell(row, especialidad_index)
    ]
    if not distribucion:
        raise ArchivoImportError("No se encontró distribución de periodos válida en la plantilla")
    return distribucion


def _parse_fechas_periodos(rows: list[list[str]]) -> list[PeriodoFechaIn]:
    if not rows:
        return []
    header, data_rows = _header_rows(rows, {"periodo", "fecha"})
    periodo_index = _buscar_columna(header, ["periodo"], default=0)
    inicio_index = _buscar_columna(
        header,
        ["fecha inicio", "fecha_inicio", "inicio", "desde", "fecha"],
        default=1,
    )
    fin_index = _buscar_columna(
        header,
        ["fecha fin", "fecha_fin", "fin", "hasta"],
        default=-1,
    )
    fechas: list[PeriodoFechaIn] = []
    for row in data_rows:
        periodo = _cell(row, periodo_index)
        if not periodo:
            continue
        fecha_inicio = _cell(row, inicio_index) or None
        fecha_fin = (_cell(row, fin_index) or None) if fin_index >= 0 else None
        fechas.append(
            PeriodoFechaIn(
                periodo=_to_int(periodo),
                fecha_inicio=fecha_inicio,
                fecha_fin=fecha_fin,
            )
        )
    return fechas


def _parse_combinaciones(rows: list[list[str]]) -> list[CombinacionIn]:
    if not rows:
        return []
    header, data_rows = _header_rows(rows, {"especialidad", "bloque"})
    esp1_index = _buscar_columna(header, ["especialidad 1", "especialidad1"], default=0)
    esp2_index = _buscar_columna(header, ["especialidad 2", "especialidad2"], default=1)
    bloque_index = _buscar_columna(header, ["bloque"], default=2)
    return [
        CombinacionIn(
            especialidad_1=_cell(row, esp1_index),
            especialidad_2=_cell(row, esp2_index),
            bloque=_cell(row, bloque_index) or "X2",
        )
        for row in data_rows
        if _cell(row, esp1_index) and _cell(row, esp2_index)
    ]


def _rows_from_values(rows: Any) -> list[list[str]]:
    table: list[list[str]] = []
    for row in rows:
        values = [_excel_cell_to_string(cell) for cell in row]
        if any(values):
            table.append(values)
    return table


def _excel_cell_to_string(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _separar_encabezado(
    rows: list[list[str]], known_columns: set[str]
) -> tuple[list[str], list[list[str]]]:
    first = [_normalizar_header(cell) for cell in rows[0]]
    has_header = any(any(alias in column for alias in known_columns) for column in first)
    if has_header:
        return first, rows[1:]
    fallback = ["id", "nombre", "semestre"]
    return fallback, rows


def _header_rows(rows: list[list[str]], known_columns: set[str]) -> tuple[list[str], list[list[str]]]:
    if not rows:
        return [], []
    first = [_normalizar_header(cell) for cell in rows[0]]
    has_header = any(any(alias in column for alias in known_columns) for column in first)
    if has_header:
        return first, rows[1:]
    return first, rows[1:]


def _key_value_rows(rows: list[list[str]]) -> dict[str, str]:
    data: dict[str, str] = {}
    if not rows:
        return data
    header = [_normalizar_header(cell) for cell in rows[0]]
    has_header = "campo" in header or "valor" in header
    data_rows = rows[1:] if has_header else rows
    key_index = _buscar_columna(header, ["campo", "parametro"], default=0) if has_header else 0
    value_index = _buscar_columna(header, ["valor"], default=1) if has_header else 1
    for row in data_rows:
        key = _normalizar_header(_cell(row, key_index))
        value = _cell(row, value_index)
        if key:
            data[key] = value
    return data


def _to_int(value: str | None) -> int:
    if not value:
        return 0
    try:
        return int(float(value))
    except ValueError as exc:
        raise ArchivoImportError(f"Valor numérico inválido: {value}") from exc


def _to_bool(value: str | None, default: bool) -> bool:
    if value is None or value == "":
        return default
    normalized = limpiar_texto(value)
    if normalized in {"TRUE", "SI", "SÍ", "1", "YES"}:
        return True
    if normalized in {"FALSE", "NO", "0"}:
        return False
    return default


def _normalizar_header(value: str) -> str:
    return limpiar_texto(value).lower()


def _buscar_columna(header: list[str], aliases: list[str], default: int) -> int:
    for index, column in enumerate(header):
        if any(column == alias or alias in column for alias in aliases):
            return index
    return default


def _period_columns(header: list[str]) -> list[tuple[int, int]]:
    columns: list[tuple[int, int]] = []
    for index, column in enumerate(header):
        if "periodo" not in column:
            continue
        digits = "".join(char for char in column if char.isdigit())
        if digits:
            columns.append((int(digits), index))
    return columns


def _scenario_column(header: list[str], period: int, period_index: int) -> int:
    for index, column in enumerate(header):
        if "escenario" in column and str(period) in column:
            return index
    return period_index + 1


def _cell(row: list[str], index: int) -> str:
    if index < 0 or index >= len(row):
        return ""
    return row[index].strip()
